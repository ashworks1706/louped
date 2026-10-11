"""reports/: decks, documents and exported figures, shown, previewed and checked."""

import asyncio
import base64
import functools
import glob
import json
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from mcp.server.mcpserver.exceptions import ToolError
from mcp_types import CallToolResult, ImageContent
from pptx import Presentation
from pptx.util import Inches
from pydantic import TypeAdapter
from test_agent import call, tools
from test_views import CLOUD, run_id  # noqa: F401  (the fixture)

from louped.analysis.views import heatmap
from louped.check import check
from louped.reports import (
    Format,
    MissingTool,
    export_figure,
    list_reports,
    number,
    outline,
    preview,
    preview_pdf,
    render,
    resolve,
)
from louped.server import create_app
from louped.stores import add_view
from louped.stores.types import PlotlyView, VegaView, View

BARS = {"kind": "vega", "title": "Caving by condition",
        "spec": {"mark": "bar", "data": {"values": [{"c": "pressure", "v": 0.9}]},
                 "encoding": {"x": {"field": "c", "type": "nominal"},
                              "y": {"field": "v", "type": "quantitative"}}}}  # fmt: skip
#: A Chrome for Kaleido: one on PATH, or the Playwright one the web tests use.
CHROME = (
    shutil.which("google-chrome")
    or shutil.which("chromium")
    or next(iter(sorted(glob.glob("/opt/pw-browsers/chromium-*/chrome-linux/chrome"))), None)
)


def deck(path: Path, run: str) -> None:
    """A deck of two slides: a number sourced in the speaker notes, and one that is not."""
    slides = Presentation()
    first = slides.slides.add_slide(slides.slide_layouts[5])
    assert first.shapes.title is not None
    first.shapes.title.text = "Caving rate 0.92"
    first.notes_slide.notes_text_frame.text = f"Source: run:{run}"  # type: ignore[union-attr]
    second = slides.slides.add_slide(slides.slide_layouts[6])
    second.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = "78%"
    slides.save(str(path))


def doc(path: Path, run: str) -> None:
    """A document whose table takes its source from the paragraph above it."""
    d = Document()
    d.add_paragraph(f"Caving by condition, run:{run}.")
    table = d.add_table(rows=2, cols=2)
    for row, (a, b) in zip(table.rows, (("condition", "caved"), ("pressure", "0.92")), strict=True):
        row.cells[0].text, row.cells[1].text = a, b
    d.add_paragraph("Unsourced: 0.5.")
    d.save(str(path))


def test_decks_and_documents_are_listed_outlined_and_checked(
    run_id: str,  # noqa: F811
    tmp_path: Path,
) -> None:
    assert list_reports() == []
    root = tmp_path / "reports"
    root.mkdir()
    deck(root / "lab.pptx", run_id)
    doc(root / "memo.docx", run_id)
    (root / "notes.txt").write_text("not shown")
    assert [(r.path, r.kind) for r in list_reports()] == [
        ("lab.pptx", "pptx"),
        ("memo.docx", "docx"),
    ]
    assert outline("lab.pptx") == [f"Caving rate 0.92\nSource: run:{run_id}", "78%"]
    assert outline("memo.docx")[1] == "condition | caved\npressure | 0.92"
    with pytest.raises(ValueError, match="louped shows"):
        outline("notes.txt")
    with pytest.raises(ValueError, match="not a name under"):
        outline("../experiments/x.md")
    got = [(i.where(), i.message.split(" has")[0]) for i in check()]
    assert got == [("reports/lab.pptx slide 2", "78%"), ("reports/memo.docx paragraph 3", "0.5")]
    # Office's lock file beside an open deck is not a report; a broken deck is an issue in it
    (root / "~$lab.pptx").write_bytes(b"lock")
    (root / "broken.docx").write_bytes(b"not a zip")
    assert "~$lab.pptx" not in [r.path for r in list_reports()]
    [broken] = [i for i in check() if i.kind == "file"]
    assert broken.file == "reports/broken.docx" and broken.message.startswith("cannot be read")


def test_a_vega_figure_exports_with_its_ref_and_trace(
    run_id: str,  # noqa: F811
    tmp_path: Path,
) -> None:
    add_view(run_id, "bars", VegaView.model_validate(BARS))
    ref = f"run:{run_id}/views/bars.json"
    formats: tuple[tuple[Format, bytes], ...] = (("svg", b"<svg"), ("png", b"\x89PNG"),
                                                 ("pdf", b"%PDF"))  # fmt: skip
    for fmt, head in formats:
        made = export_figure(ref, fmt, name="bars")
        assert (made.path, made.kind, made.ref) == (f"figures/bars.{fmt}", fmt, ref)
        assert (tmp_path / "reports" / made.path).read_bytes().startswith(head)
    assert [(r.path, r.ref) for r in list_reports()] == [
        (f"figures/bars.{fmt}", ref) for fmt in ("pdf", "png", "svg")
    ]
    sidecar_file = tmp_path / "reports" / "figures" / "bars.svg.refs.json"
    sidecar = json.loads(sidecar_file.read_text())
    assert sidecar["ref"] == ref and sidecar["trace"]["steps"][-1]["what"] == "run"
    with pytest.raises(ValueError, match="names no figure"):
        export_figure(f"run:{run_id}")
    with pytest.raises(ValueError, match="not a file name"):
        export_figure(ref, name="../x")
    with pytest.raises(FileExistsError, match="exists"):
        export_figure(ref, "svg", name="bars")
    assert export_figure(ref, "svg", name="bars", replace=True).path == "figures/bars.svg"
    assert check() == []
    # an exported figure whose run is gone no longer checks
    sidecar["ref"] = "run:m-gone/views/bars.json"
    sidecar_file.write_text(json.dumps(sidecar))
    [gone] = check()
    assert gone.file == "reports/figures/bars.svg.refs.json" and "does not resolve" in gone.message
    # a refs file that holds no ref is named, by the listing and by the check
    sidecar_file.write_text("{")
    with pytest.raises(ValueError, match=r"bars\.svg\.refs\.json is not a figure's refs file"):
        list_reports()
    [broken] = check()
    assert broken.kind == "file" and "refs file" in broken.message


def test_a_plotly_figure_exports_through_kaleido_in_chrome(
    run_id: str,  # noqa: F811
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    add_view(run_id, "cloud", PlotlyView.model_validate(CLOUD))
    ref = f"run:{run_id}/views/cloud.json"
    monkeypatch.delenv("BROWSER_PATH", raising=False)
    monkeypatch.setenv("PATH", str(tmp_path))  # no Chrome on it
    with pytest.raises(MissingTool, match="through Chrome"):
        export_figure(ref, "png")
    if CHROME is None:
        pytest.skip("no Chrome to export through")
    monkeypatch.setenv("BROWSER_PATH", CHROME)
    made = export_figure(ref, "png")
    assert made.path == f"figures/{run_id}-cloud.png"
    assert (tmp_path / "reports" / made.path).read_bytes().startswith(b"\x89PNG")


@functools.cache
def converts() -> bool:
    """Whether a LibreOffice here can make a PDF: some installs carry only its core."""
    office = shutil.which("soffice")
    if office is None:
        return False
    with tempfile.TemporaryDirectory() as folder:
        (Path(folder) / "t.txt").write_text("x")
        convert = [office, "--headless", "--convert-to", "pdf", "--outdir", folder]
        subprocess.run([*convert, f"{folder}/t.txt"], capture_output=True, timeout=120, check=False)
        return (Path(folder) / "t.pdf").is_file()


def test_a_conversion_that_hangs_is_an_error(
    run_id: str,  # noqa: F811
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    (tmp_path / "reports").mkdir()
    deck(tmp_path / "reports" / "lab.pptx", run_id)
    (tmp_path / "bin").mkdir()
    (tmp_path / "bin" / "soffice").write_text("#!/bin/sh\nsleep 1\n")
    (tmp_path / "bin" / "soffice").chmod(0o755)
    monkeypatch.setenv("PATH", str(tmp_path / "bin"))

    def hangs(cmd: list[str], **kwargs: object) -> None:
        raise subprocess.TimeoutExpired(cmd, 180)

    monkeypatch.setattr(subprocess, "run", hangs)
    with pytest.raises(RuntimeError, match="took over 180s"):
        preview_pdf("lab.pptx")


def test_a_deck_previews_as_a_pdf_through_libreoffice(
    run_id: str,  # noqa: F811
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    if not converts():
        pytest.skip("no LibreOffice that converts")
    (tmp_path / "reports").mkdir()
    deck(tmp_path / "reports" / "lab.pptx", run_id)
    pdf = preview_pdf("lab.pptx")
    assert pdf.read_bytes().startswith(b"%PDF") and preview_pdf("lab.pptx") == pdf  # kept
    monkeypatch.setenv("PATH", str(tmp_path))  # no LibreOffice on it
    (tmp_path / "reports" / "lab.pptx").write_bytes(b"changed")
    with pytest.raises(MissingTool, match="LibreOffice is not installed"):
        preview_pdf("lab.pptx")


def test_reports_through_the_api_and_mcp(
    run_id: str,  # noqa: F811
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    add_view(run_id, "bars", VegaView.model_validate(BARS))
    mcp = tools()
    made = call(mcp, "export_figure", ref=f"run:{run_id}/views/bars.json")
    assert made["path"] == f"figures/{run_id}-bars.svg"
    assert [r["path"] for r in call(mcp, "reports")] == [made["path"]]
    api = TestClient(create_app(launching=True), base_url="http://localhost")
    assert api.get("/api/reports/file", params={"path": made["path"]}).text.startswith("<svg")
    assert api.get("/api/reports/file", params={"path": "figures/none.svg"}).status_code == 404
    assert api.get("/api/reports/outline", params={"path": made["path"]}).status_code == 400
    deck(tmp_path / "reports" / "lab.pptx", run_id)
    monkeypatch.setenv("PATH", str(tmp_path))
    missing = api.get("/api/reports/preview", params={"path": "lab.pptx"})
    assert missing.status_code == 501 and "LibreOffice" in missing.json()["detail"]
    nope = api.post("/api/reports/figures", json={"ref": f"run:{run_id}/views/none.json"})
    assert nope.status_code == 404
    again = api.post("/api/reports/figures", json={"ref": f"run:{run_id}/views/bars.json"})
    assert again.status_code == 409
    readonly = TestClient(create_app(launching=False), base_url="http://localhost")
    assert readonly.post("/api/reports/figures", json={"ref": "x"}).status_code == 403


def png(result: object) -> bytes:
    """The image a tool returned."""
    assert isinstance(result, CallToolResult) and not result.is_error, result
    [image] = result.content
    assert isinstance(image, ImageContent) and image.mime_type == "image/png"
    return base64.b64decode(image.data)


def test_the_agent_sees_a_figure_as_a_png_before_and_after_it_keeps_it(
    run_id: str,  # noqa: F811
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    mcp = tools()
    assert png(asyncio.run(mcp.call_tool("preview_view", {"view": BARS}))).startswith(b"\x89PNG")
    add_view(run_id, "bars", VegaView.model_validate(BARS))
    saved = asyncio.run(mcp.call_tool("preview_view", {"ref": f"run:{run_id}/views/bars.json"}))
    assert png(saved).startswith(b"\x89PNG")
    with pytest.raises(ToolError, match="give view or ref, one of them"):
        asyncio.run(mcp.call_tool("preview_view", {}))
    # a view that does not check says what is wrong in it, so the agent can fix the spec
    wrong = r"louped answered 400: the view does not check: vega\.spec: Input should be a valid"
    with pytest.raises(ToolError, match=wrong + " dictionary$"):
        asyncio.run(mcp.call_tool("preview_view", {"view": {**BARS, "spec": "bar"}}))
    with pytest.raises(ToolError, match="plotly: Value error, animation autoplay and loop play"):
        moves = {**CLOUD, "animation": {"autoplay": True}}
        asyncio.run(mcp.call_tool("preview_view", {"view": moves}))
    api = TestClient(create_app(launching=True), base_url="http://localhost")
    drawn = api.post("/api/views/preview", json={"view": BARS})
    assert drawn.headers["content-type"] == "image/png" and drawn.content.startswith(b"\x89PNG")
    lens = heatmap("lens", [[0.5]], ["0"], ["0"], "position", "layer")
    for body, code, why in (
        ({"view": {**BARS, "spec": {"mark": "nope", "data": {"values": []}}}}, 400, "failed"),
        ({"view": lens}, 400, "a heatmap figure, which only the app draws"),
        ({"ref": f"run:{run_id}/views/none.json"}, 404, "none.json"),
        ({"view": BARS, "ref": f"run:{run_id}/views/bars.json"}, 400, "one of them"),
    ):
        answer = api.post("/api/views/preview", json=body)
        assert answer.status_code == code and why in answer.json()["detail"], answer.text
    readonly = TestClient(create_app(launching=False), base_url="http://localhost")
    assert readonly.post("/api/views/preview", json={"view": BARS}).status_code == 403
    # a plotly figure draws in Chrome; without it, the answer says so
    surface = {**CLOUD, "data": [{"type": "surface", "z": [[0, 1], [1, 0]]}],
               "animation": {"orbit": True}}  # fmt: skip
    monkeypatch.delenv("BROWSER_PATH", raising=False)
    monkeypatch.setenv("PATH", str(tmp_path))  # no Chrome on it
    with pytest.raises(
        ToolError, match="louped answered 501: Plotly exports figures through Chrome"
    ):
        asyncio.run(mcp.call_tool("preview_view", {"view": surface}))
    if CHROME is None:
        pytest.skip("no Chrome to draw through")
    monkeypatch.setenv("BROWSER_PATH", CHROME)
    assert png(asyncio.run(mcp.call_tool("preview_view", {"view": surface}))).startswith(b"\x89PNG")


def test_the_figure_recipes_agents_copy_check_and_vega_ones_draw() -> None:
    import louped

    skill = Path(louped.__file__).parent / "templates" / "agent" / "skills" / "change-ui"
    recipes = re.findall(r"```json\n(.*?)```", (skill / "figures.md").read_text(), re.S)
    kinds = [TypeAdapter(View).validate_json(r).kind for r in recipes]
    assert kinds == ["plotly", "plotly", "plotly", "vega", "vega"]
    for recipe in recipes[3:]:
        assert preview(json.loads(recipe)).startswith(b"\x89PNG")
    assert "(figures.md)" in (skill / "SKILL.md").read_text()


@pytest.fixture
def caving() -> str:
    """A run under experiment 1b with a metric logged twice and a figure."""
    import mlflow

    from louped.tracking import log_json, start_run

    with start_run("1b", name="caving") as r:
        mlflow.log_metric("pressure/r_uy", 0.5, step=0)
        mlflow.log_metric("pressure/r_uy", 0.91234, step=1)
        mlflow.log_metric("n", 40)
        log_json(BARS, "views/bars.json")
    return f"m-{r.info.run_id}"


def test_live_refs_resolve_to_latest_values_and_figures_or_say_why(caving: str) -> None:
    run = f"run:{caving}"
    got = {r.text: r for r in resolve([
        f"{run} pressure/r_uy", f"{run} pressure/r_uy :.1%", f"{run} n",
        f"{run}/views/bars.json", f"{run} nope", "run:m-404 n", f"{run} n :q",
        f"{run}/views/bars.json :.1%", f"{run}/views/none.json", f"{run} a b",
        "experiment:1b n",
    ])}  # fmt: skip
    assert [got[f"{run} {m}"].value for m in ("pressure/r_uy", "pressure/r_uy :.1%", "n")] == [
        "0.912",
        "91.2%",
        "40",
    ]
    figure = got[f"{run}/views/bars.json"]
    assert figure.view is not None and figure.view.title == "Caving by condition"
    assert figure.error is None and figure.value is None
    errors = {t: r.error or "" for t, r in got.items() if r.error}
    assert len(errors) == 7
    assert errors[f"{run} nope"].startswith(f"run {caving} has no metric nope; it has n, ")
    assert errors["run:m-404 n"] == "not found: m-404"
    assert "'q' is not a format" in errors[f"{run} n :q"]
    assert "a format is for a metric" in errors[f"{run}/views/bars.json :.1%"]
    assert "not found" in errors[f"{run}/views/none.json"]
    assert "a live ref is" in errors[f"{run} a b"]
    assert "is not a run" in errors["experiment:1b n"]


def test_numbers_keep_three_significant_decimals_unless_given_a_format() -> None:
    assert [number(v) for v in (0.91234, 0.0912345, 1234.5678, 3.0, 0.000123456, -0.5)] == [
        "0.912",
        "0.0912",
        "1234.568",
        "3",
        "0.000123",
        "-0.5",
    ]
    assert (number(0.873, ".1%"), number(0.873, ".2f"), number(float("nan"))) == (
        "87.3%",
        "0.87",
        "nan",
    )


def test_live_refs_render_check_and_resolve_over_the_api(caving: str, tmp_path: Path) -> None:
    text = (
        f"Caving rose to {{{{run:{caving} pressure/r_uy :.0%}}}} from 0.5.\n\n"
        f"{{{{run:{caving}/views/bars.json}}}}\n\n"
        f"Broken: {{{{run:{caving} gone}}}} and 0.3.\n"
    )
    rendered, broken = render(text)
    assert rendered.splitlines()[0] == "Caving rose to 91% from 0.5."
    assert f"{{{{run:{caving}/views/bars.json}}}}" in rendered  # a figure stays as written
    assert [b.text for b in broken] == [f"run:{caving} gone"]
    (tmp_path / "reports").mkdir()
    (tmp_path / "reports" / "1b.md").write_text(text)
    # a live ref sources its paragraph; one that does not resolve is one ref issue, not two
    issues = [(i.line, i.kind, i.message.split(" does")[0]) for i in check()]
    assert issues == [(5, "ref", f"{{{{run:{caving} gone}}}}")]
    api = TestClient(create_app(launching=False), base_url="http://localhost")
    got = api.get("/api/live", params={"ref": [f"run:{caving} n", "run:m-404 n"]}).json()
    assert [(g["value"], g["error"]) for g in got] == [("40", None), (None, "not found: m-404")]


def test_reports_group_by_experiment_from_front_matter_folder_or_refs(
    caving: str, tmp_path: Path
) -> None:
    from test_stores import write_experiment

    for name in ("1b", "1c"):
        write_experiment(tmp_path, name, "honesty", "active")
    root = tmp_path / "reports"
    (root / "1c").mkdir(parents=True)
    (root / "paper-notes").mkdir()
    (root / "front.md").write_text(f"---\nexperiment: 1c\n---\n\nSee run:{caving}.\n")
    (root / "1c" / "notes.md").write_text(f"Cites run:{caving} and run:{caving}/raw/x.jsonl.\n")
    (root / "paper-notes" / "related.md").write_text("No refs here.\n")
    (root / "mixed.md").write_text(f"run:{caving} and experiment:1c/views/a.json\n")
    (root / "agree.md").write_text(f"{{{{run:{caving} n}}}}, experiment:1b/views/a.json.\n")
    deck(root / "lab.pptx", caving)
    figure = export_figure(f"run:{caving}/views/bars.json").path
    got = {r.path: (r.experiment, r.runs) for r in list_reports()}
    assert got == {
        "agree.md": ("1b", [caving]),
        figure: ("1b", [caving]),
        "front.md": ("1c", [caving]),
        "1c/notes.md": ("1c", [caving]),
        "lab.pptx": ("1b", [caving]),
        "mixed.md": (None, [caving]),
        "paper-notes/related.md": (None, []),
    }
    # read again once it changes
    (root / "paper-notes" / "related.md").write_text(f"Now run:{caving}, a longer text.\n")
    assert [r.experiment for r in list_reports() if r.path.startswith("paper")] == ["1b"]
