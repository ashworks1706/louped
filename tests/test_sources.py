"""The project's sources/: files and URLs kept, read page by page, searched and quoted."""

import json
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import httpx
import nbformat
import pytest
from docx import Document
from fastapi.testclient import TestClient
from pptx import Presentation
from test_agent import call, tools

from louped.server import create_app
from louped.sources import add_pin, add_source, list_sources, page_text, pins, search


def pdf(*pages: str) -> bytes:
    """A small PDF with one line of text on each page."""
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", ""]
    kids = []
    font = 3 + 2 * len(pages)
    for i, text in enumerate(pages):
        page, content = 3 + 2 * i, 4 + 2 * i
        kids.append(f"{page} 0 R")
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET"
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content} 0 R"
                    f" /Resources << /Font << /F1 {font} 0 R >> >> >>")  # fmt: skip
        objs.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
    objs[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(pages)} >>"
    objs.append("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out, offsets = b"%PDF-1.4\n", []
    for n, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{n} 0 obj\n{body}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{o:010d} 00000 n \n" for o in offsets).encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


def test_files_of_every_kind_are_kept_read_by_page_and_searched(tmp_path: Path) -> None:
    (tmp_path / "paper.pdf").write_bytes(pdf("Models cave to pushback", "Evidence helps them"))
    deck = Presentation()
    for words in ("Lab meeting", "Caving rate by condition"):
        title = deck.slides.add_slide(deck.slide_layouts[5]).shapes.title
        assert title is not None
        title.text = words
    deck.save(str(tmp_path / "deck.pptx"))
    doc = Document()
    doc.add_paragraph("Notes on sycophancy under pressure.")
    doc.save(str(tmp_path / "notes.docx"))
    nb = nbformat.v4.new_notebook(cells=[nbformat.v4.new_markdown_cell("# Embeddings"),
                                         nbformat.v4.new_code_cell("umap.fit(x)")])  # fmt: skip
    nbformat.write(nb, str(tmp_path / "explore.ipynb"))
    (tmp_path / "plan.md").write_text("# The 1B plan\n\nAnti-pressure DPO.\n")

    paper = add_source(str(tmp_path / "paper.pdf"), title="Sycophancy under pushback")
    assert (paper.key, paper.kind, paper.pages, paper.origin) == (
        "sycophancy-under-pushback", "pdf", 2, None,
    )  # fmt: skip
    assert add_source(str(tmp_path / "deck.pptx")).pages == 2
    assert add_source(str(tmp_path / "notes.docx")).pages == 1
    assert add_source(str(tmp_path / "explore.ipynb")).pages == 2
    assert add_source(str(tmp_path / "plan.md")).title == "The 1B plan"
    assert [s.key for s in list_sources()] == [
        "deck", "explore", "notes", "plan", "sycophancy-under-pushback",
    ]  # fmt: skip
    kept = tmp_path / "sources"
    assert (kept / "sycophancy-under-pushback.pdf").read_bytes().startswith(b"%PDF")
    index = json.loads((kept / "index.json").read_text())
    assert index[4]["sha256"] == paper.sha256

    assert page_text("sycophancy-under-pushback", 2).strip() == "Evidence helps them"
    assert page_text("deck", 2) == "Caving rate by condition"
    [hit] = search("pushback caves")  # every word, stemmed: "caves" finds "cave"
    assert (hit.key, hit.page, hit.title) == ("sycophancy-under-pushback", 1,
                                              "Sycophancy under pushback")  # fmt: skip
    assert "[pushback]" in hit.snippet
    assert [h.key for h in search("sycophancy")] == ["notes"]
    assert search('"; DROP TABLE pages; --') == []  # words only, never FTS syntax
    # a file changed on disk is read again before the next search
    (kept / "plan.md").write_text("# The 1B plan\n\nNow with a reward model.\n")
    index[3]["sha256"] = "changed"
    (kept / "index.json").write_text(json.dumps(index))
    assert [h.key for h in search("reward model")] == ["plan"]

    # the same file again returns the one kept; another file under a taken key is refused
    assert add_source(str(tmp_path / "deck.pptx")).key == "deck"
    with pytest.raises(ValueError, match="already a source"):
        add_source(str(tmp_path / "plan.md"), key="deck")
    with pytest.raises(ValueError, match="not a source key"):
        add_source(str(tmp_path / "paper.pdf"), key="Bad Key")
    (tmp_path / "data.csv").write_text("a,b\n")
    with pytest.raises(ValueError, match="louped reads"):
        add_source(str(tmp_path / "data.csv"))
    (tmp_path / "broken.pdf").write_bytes(b"not a pdf")
    with pytest.raises(ValueError, match="not a readable pdf"):
        add_source(str(tmp_path / "broken.pdf"))
    assert not (kept / "broken.pdf.tmp").exists() and not (kept / "broken.pdf").exists()
    with pytest.raises(ValueError, match="no page 3"):
        page_text("deck", 3)


ATOM = """<feed xmlns="http://www.w3.org/2005/Atom"><entry><title>Towards Understanding
  Sycophancy in Language Models</title></entry></feed>"""


def web(request: httpx.Request) -> httpx.Response:
    url = str(request.url)
    if url.startswith("https://export.arxiv.org/api/query"):
        assert request.url.params["id_list"] == "2310.13548"
        return httpx.Response(200, text=ATOM)
    if url == "https://arxiv.org/pdf/2310.13548":
        return httpx.Response(200, content=pdf("Sycophancy is a general behavior"),
                              headers={"content-type": "application/pdf"})  # fmt: skip
    if url == "https://docs.google.com/document/d/abc_1/export?format=pdf":
        return httpx.Response(200, content=pdf("Shared draft"))
    if url == "https://raw.githubusercontent.com/o/r/main/nb/explore.ipynb":
        nb = nbformat.v4.new_notebook(cells=[nbformat.v4.new_code_cell("print(1)")])
        return httpx.Response(200, text=nbformat.writes(nb))
    if url == "https://example.org/to-local.pdf":
        return httpx.Response(302, headers={"location": "https://10.0.0.1/x.pdf"})
    if url == "https://example.org/blog":
        return httpx.Response(200, text="<html></html>", headers={"content-type": "text/html"})
    return httpx.Response(404)


def test_urls_are_fetched_once_with_where_they_came_from() -> None:
    to = httpx.MockTransport(web)
    paper = add_source("https://arxiv.org/abs/2310.13548v4", transport=to)
    assert (paper.key, paper.title, paper.origin) == (
        "towards-understanding-sycophancy-in-language-mod",
        "Towards Understanding Sycophancy in Language Models",
        "https://arxiv.org/abs/2310.13548v4",
    )
    assert page_text(paper.key, 1).strip() == "Sycophancy is a general behavior"
    doc = add_source("https://docs.google.com/document/d/abc_1/edit", key="draft", transport=to)
    assert (doc.kind, doc.origin) == ("pdf", "https://docs.google.com/document/d/abc_1/edit")
    nb = add_source("https://colab.research.google.com/github/o/r/blob/main/nb/explore.ipynb",
                    transport=to)  # fmt: skip
    assert (nb.key, nb.kind, nb.pages) == ("explore", "ipynb", 1)
    for url, why in (
        ("http://arxiv.org/abs/2310.13548", "https only"),
        ("https://colab.research.google.com/drive/1xyz", "Download .ipynb"),
        ("https://example.org/blog", "a web page, not a file"),
        ("https://127.0.0.1:8000/x.pdf", "not on the internet"),
        ("https://[::1]/x.pdf", "not on the internet"),
        ("https://localhost/x.pdf", "not on the internet"),
        ("https://example.org/to-local.pdf", "not on the internet"),
    ):
        with pytest.raises(ValueError, match=why):
            add_source(url, transport=to)
    with pytest.raises(httpx.HTTPStatusError):
        add_source("https://example.org/missing.pdf", transport=to)


def test_sources_through_the_api_and_mcp(tmp_path: Path) -> None:
    (tmp_path / "paper.pdf").write_bytes(pdf("Models cave to pushback"))
    mcp = tools()
    added = call(mcp, "add_source", location=str(tmp_path / "paper.pdf"), key="pushback")
    assert added["key"] == "pushback"
    assert [s["key"] for s in call(mcp, "sources")] == ["pushback"]
    [hit] = call(mcp, "search_sources", query="pushback")
    assert (hit["key"], hit["page"]) == ("pushback", 1)
    assert call(mcp, "source_page", key="pushback", page=1)["text"].strip() == (
        "Models cave to pushback"
    )
    api = TestClient(create_app(launching=True), base_url="http://localhost")
    assert api.get("/api/sources/nope/pages/1").status_code == 404
    assert api.get("/api/sources/pushback/pages/9").status_code == 400
    bad = api.post("/api/sources", json={"location": str(tmp_path / "none.pdf")})
    assert bad.status_code == 400
    outside = api.post("/api/sources", json={"location": "/etc/hostname"})
    assert outside.status_code == 400 and "not a name under" in outside.json()["detail"]
    (tmp_path / "notes.md").write_text("# Notes\n")
    assert api.post("/api/sources", json={"location": "notes.md"}).json()["key"] == "notes"
    readonly = TestClient(create_app(launching=False), base_url="http://localhost")
    assert readonly.post("/api/sources", json={"location": "x"}).status_code == 403
    assert readonly.get("/api/sources/search", params={"q": "cave"}).json()[0]["key"] == "pushback"


def test_a_pin_holds_words_on_its_page_and_is_kept_once(tmp_path: Path) -> None:
    (tmp_path / "paper.pdf").write_bytes(pdf("Models cave to pushback when pressed", "Evidence"))
    mcp = tools()
    call(mcp, "add_source", location=str(tmp_path / "paper.pdf"), key="pushback")
    got = call(mcp, "pin", key="pushback", page=1, quote="cave  to\npushback",
               note="the caving claim", links=["experiment:pressure"])  # fmt: skip
    assert (got["exact"], got["start"]) == ("cave to pushback", 6)
    # the same words again: one pin, the note kept, the links added to
    again = add_pin("pushback", 1, "cave to pushback", links=["run:m-1"])
    assert (again.id, again.note) == (got["id"], "the caving claim")
    assert again.links == ["experiment:pressure", "run:m-1"]
    add_pin("pushback", 2, "Evidence")
    assert [(p.page, p.exact) for p in pins("pushback")] == [(1, "cave to pushback"),
                                                             (2, "Evidence")]  # fmt: skip
    with pytest.raises(ValueError, match="is not on page 2"):
        add_pin("pushback", 2, "cave to pushback")
    with pytest.raises(KeyError):
        add_pin("nope", 1, "x")
    api = TestClient(create_app(launching=True), base_url="http://localhost")
    assert api.get("/api/sources/pushback/file").content.startswith(b"%PDF")
    assert api.get("/api/sources/nope/file").status_code == 404
    paraphrase = api.post("/api/pins", json={"key": "pushback", "page": 1, "quote": "caves"})
    assert paraphrase.status_code == 400
    assert api.post("/api/pins", json={"key": "nope", "page": 1, "quote": "x"}).status_code == 404
    assert [p["page"] for p in call(mcp, "pins", key="pushback")] == [1, 2]
    gone = {"content-type": "application/json"}
    assert api.request("DELETE", f"/api/pins/{got['id']}", headers=gone).status_code == 200
    assert api.request("DELETE", f"/api/pins/{got['id']}", headers=gone).status_code == 404
    assert [p.page for p in pins()] == [2]
    readonly = TestClient(create_app(launching=False), base_url="http://localhost")
    assert readonly.post("/api/pins", json={"key": "pushback", "page": 2,
                                            "quote": "Evidence"}).status_code == 403  # fmt: skip


def test_a_quote_selected_across_lines_is_found_and_pins_written_at_once_are_all_kept(
    tmp_path: Path,
) -> None:
    (tmp_path / "notes.md").write_text("Models cave to\npushback when pressed.\n")
    add_source(str(tmp_path / "notes.md"))
    # a PDF viewer's selection can lose the break between lines, or gain one
    assert add_pin("notes", 1, "cave topushback").exact == "cave to pushback"
    assert add_pin("notes", 1, "cave to pushback").id == pins("notes")[0].id
    words = ["Models", "cave", "to", "pushback", "when", "pressed."]
    with ThreadPoolExecutor(8) as pool:
        list(pool.map(lambda w: add_pin("notes", 1, w), words))
    assert len(pins("notes")) == 1 + len(words)
    # an index entry naming a file outside sources/ is refused, not read
    index = tmp_path / "sources" / "index.json"
    entries = json.loads(index.read_text())
    entries[0]["file"] = "../notes.md"
    index.write_text(json.dumps(entries))
    api = TestClient(create_app(), base_url="http://localhost")
    assert api.get("/api/sources/notes/file").status_code == 400
