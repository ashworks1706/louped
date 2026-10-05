"""The project's sources: the papers, docs, slides and notebooks its claims rest on, kept in
sources/ beside experiments/ and committed with them, so a claim can quote and cite them.

sources/index.json lists each one:
- key: what a citation names it by;
- title and kind (pdf, pptx, docx, ipynb, md, txt);
- file, under sources/;
- origin: the URL it came from, or None for a file you added;
- the sha256 of its file.

Its text is read page by page into a search index at <home>/sources.db, rebuilt from the files
whenever one changes, so the index is never committed. A page is a PDF page, a slide, a notebook
cell, or the whole of a Markdown, text or Word file (Word files have no fixed pages).

A URL is fetched once and kept. An arXiv page becomes its PDF and the paper's title. A shared
Google Doc or Slides deck becomes a PDF snapshot, because a quote must point at text that cannot
change under it. A notebook on GitHub, or opened from GitHub in Colab, is fetched raw.
"""

from __future__ import annotations

import hashlib
import json
import os
import re
import sqlite3
import xml.etree.ElementTree as ET
from contextlib import closing
from datetime import UTC, datetime
from pathlib import Path
from typing import Literal

import httpx
from pydantic import BaseModel, TypeAdapter

from louped.core.paths import NAME, experiments_dir, home, inside

Kind = Literal["pdf", "pptx", "docx", "ipynb", "md", "txt"]
KINDS: dict[str, Kind] = {".pdf": "pdf", ".pptx": "pptx", ".docx": "docx", ".ipynb": "ipynb",
                          ".md": "md", ".markdown": "md", ".txt": "txt"}  # fmt: skip
#: The largest file a URL may bring in.
MAX_BYTES = 100 * 2**20
INDEX = "index.json"


class Source(BaseModel):
    key: str
    title: str
    kind: Kind
    #: Its file, relative to sources/.
    file: str
    #: The URL it was fetched from; None for a file added from disk.
    origin: str | None = None
    sha256: str
    #: When it was added (UTC).
    added: datetime
    pages: int


class Hit(BaseModel):
    key: str
    title: str
    page: int
    #: The matching words in [brackets], with a little text around them.
    snippet: str


_LIST: TypeAdapter[list[Source]] = TypeAdapter(list[Source])


def sources_dir() -> Path:
    """sources/ at the project's root, beside experiments/."""
    return experiments_dir().parent / "sources"


def list_sources() -> list[Source]:
    """Every source in sources/index.json, by key; none when there is no index."""
    path = sources_dir() / INDEX
    if not path.is_file():
        return []
    return _LIST.validate_json(path.read_bytes())


def source(key: str) -> Source:
    """One source by key; KeyError naming the keys there are when it is not one."""
    found = {s.key: s for s in list_sources()}
    if key not in found:
        raise KeyError(f"no source {key!r}; there are {sorted(found)}")
    return found[key]


def add_source(location: str, key: str | None = None, title: str | None = None,
               transport: httpx.BaseTransport | None = None, project_only: bool = False,
               ) -> Source:  # fmt: skip
    """Add a file or a URL to sources/ and its text to the search index. A key already taken is
    refused unless it is the same file again, which returns the one there. With project_only (the
    server's way), a file must be inside the project: what the app is asked for over HTTP never
    reads elsewhere on the machine."""
    root = sources_dir()
    if re.match(r"^https?://", location):
        if not location.startswith("https://"):
            raise ValueError(f"{location} is not https: louped fetches over https only")
        url, found_title, suggested = _resolve(location, transport)
        data, kind = _fetch(url, transport)
        origin: str | None = location
    else:
        # a path relative to the project, or an absolute one inside it
        path = (inside(experiments_dir().parent, location) if project_only
                else Path(location).expanduser())  # fmt: skip
        if not path.is_file():
            raise FileNotFoundError(f"no file {location}")
        kind = _kind(path.name)
        data, origin, found_title, suggested = path.read_bytes(), None, None, path.stem
    key = key or _slug(title or found_title or suggested)
    if not NAME.match(key):
        raise ValueError(f"{key!r} is not a source key: lowercase letters, digits and -")
    sha = hashlib.sha256(data).hexdigest()
    have = {s.key: s for s in list_sources()}
    if key in have:
        if have[key].sha256 == sha:
            return have[key]
        raise ValueError(f"{key} is already a source ({have[key].file}): give another key")
    target = inside(root, f"{key}.{'md' if kind == 'md' else kind}")
    root.mkdir(exist_ok=True)
    tmp = target.with_suffix(target.suffix + ".tmp")
    tmp.write_bytes(data)
    try:
        text, own_title = _read(tmp, kind)
    except Exception as exc:
        tmp.unlink()
        raise ValueError(f"{location} is not a readable {kind}: {exc}") from exc
    os.replace(tmp, target)
    added = Source(key=key, title=title or found_title or own_title or suggested, kind=kind,
                   file=target.name, origin=origin, sha256=sha, added=datetime.now(UTC),
                   pages=len(text))  # fmt: skip
    _write_index([*have.values(), added])
    return added


def page_text(key: str, page: int) -> str:
    """One page of a source's text, counted from 1."""
    found = source(key)
    pages, _ = _read(inside(sources_dir(), found.file), found.kind)
    if not 1 <= page <= len(pages):
        raise ValueError(f"{key} has pages 1 to {len(pages)}; there is no page {page}")
    return pages[page - 1]


def search(query: str, limit: int = 10) -> list[Hit]:
    """The pages that hold every word of query, best first."""
    words = re.findall(r"\w+", query)
    if not words:
        return []
    match = " ".join(f'"{w}"' for w in words)
    titles = {s.key: s.title for s in list_sources()}
    with closing(_index()) as db:
        rows = db.execute(
            "SELECT key, page, snippet(pages, 2, '[', ']', '…', 16) FROM pages "
            "WHERE pages MATCH ? ORDER BY rank LIMIT ?",
            (match, limit),
        ).fetchall()
    return [Hit(key=k, title=str(titles.get(k, k)), page=p, snippet=s) for k, p, s in rows]


def _index() -> sqlite3.Connection:
    """The search index, brought up to date with sources/: a changed file is read again, a
    removed one dropped."""
    home().mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(home() / "sources.db")
    db.execute("CREATE TABLE IF NOT EXISTS files (key TEXT PRIMARY KEY, sha256 TEXT)")
    db.execute("CREATE VIRTUAL TABLE IF NOT EXISTS pages USING fts5(key UNINDEXED, "
               "page UNINDEXED, text, tokenize='porter unicode61')")  # fmt: skip
    have = dict(db.execute("SELECT key, sha256 FROM files").fetchall())
    now = {s.key: s for s in list_sources()}
    with db:
        for gone in have.keys() - now.keys():
            db.execute("DELETE FROM pages WHERE key = ?", (gone,))
            db.execute("DELETE FROM files WHERE key = ?", (gone,))
        for key, s in now.items():
            path = inside(sources_dir(), s.file)
            if have.get(key) == s.sha256 or not path.is_file():
                continue
            pages, _ = _read(path, s.kind)
            db.execute("DELETE FROM pages WHERE key = ?", (key,))
            db.executemany("INSERT INTO pages (key, page, text) VALUES (?, ?, ?)",
                           [(key, i, t) for i, t in enumerate(pages, 1)])  # fmt: skip
            db.execute("INSERT OR REPLACE INTO files VALUES (?, ?)", (key, s.sha256))
    return db


def _write_index(found: list[Source]) -> None:
    path = sources_dir() / INDEX
    tmp = path.with_suffix(".tmp")
    data = _LIST.dump_python(sorted(found, key=lambda s: s.key), mode="json")
    tmp.write_text(json.dumps(data, indent=2) + "\n", encoding="utf-8")
    os.replace(tmp, path)


def _kind(name: str) -> Kind:
    suffix = Path(name).suffix.lower()
    if suffix not in KINDS:
        raise ValueError(f"{name}: louped reads {', '.join(sorted(KINDS))} files")
    return KINDS[suffix]


def _slug(text: str) -> str:
    return re.sub(r"-+", "-", re.sub(r"[^a-z0-9]+", "-", text.lower())).strip("-")[:48] or "source"


_ARXIV = re.compile(r"^https://(?:www\.)?arxiv\.org/(?:abs|pdf)/([0-9]{4}\.[0-9]{4,5})(?:v\d+)?")
_GOOGLE = re.compile(r"^https://docs\.google\.com/(document|presentation)/d/([\w-]+)")
_GITHUB = re.compile(r"^https://(?:colab\.research\.google\.com/github|github\.com)/"
                     r"([^/]+)/([^/]+)/blob/(.+)$")  # fmt: skip


def _resolve(url: str, transport: httpx.BaseTransport | None) -> tuple[str, str | None, str]:
    """What to fetch for a URL, the title when the place says it, and a key to suggest."""
    if m := _ARXIV.match(url):
        paper = m.group(1)
        return f"https://arxiv.org/pdf/{paper}", _arxiv_title(paper, transport), f"arxiv-{paper}"
    if m := _GOOGLE.match(url):
        doc, ident = m.groups()
        export = "export?format=pdf" if doc == "document" else "export/pdf"
        return f"https://docs.google.com/{doc}/d/{ident}/{export}", None, f"google-{doc}"
    if m := _GITHUB.match(url):
        owner, repo, rest = m.groups()
        return (f"https://raw.githubusercontent.com/{owner}/{repo}/{rest}", None,
                Path(rest).stem)  # fmt: skip
    if url.startswith("https://colab.research.google.com/"):
        raise ValueError("a Colab notebook kept in Google Drive needs your sign-in: download it "
                         "(File > Download .ipynb) and add the file")  # fmt: skip
    return url, None, Path(httpx.URL(url).path).stem or "source"


def _arxiv_title(paper: str, transport: httpx.BaseTransport | None) -> str | None:
    """The paper's title from arXiv's API."""
    with httpx.Client(transport=transport, timeout=30) as client:
        r = client.get("https://export.arxiv.org/api/query", params={"id_list": paper})
    r.raise_for_status()
    atom = "{http://www.w3.org/2005/Atom}"
    entry = ET.fromstring(r.text).find(f"{atom}entry")
    title = entry.findtext(f"{atom}title") if entry is not None else None
    return " ".join(title.split()) if title else None


def _fetch(url: str, transport: httpx.BaseTransport | None) -> tuple[bytes, Kind]:
    """A URL's bytes, at most MAX_BYTES, and their kind from the type the server gives."""
    with (httpx.Client(transport=transport, timeout=60, follow_redirects=True) as client,
          client.stream("GET", url) as r):  # fmt: skip
        r.raise_for_status()
        if not str(r.url).startswith("https://"):
            raise ValueError(f"{url} redirected to {r.url}, which is not https")
        chunks, size = [], 0
        for chunk in r.iter_bytes():
            size += len(chunk)
            if size > MAX_BYTES:
                raise ValueError(f"{url} is over {MAX_BYTES // 2**20} MB")
            chunks.append(chunk)
        kind_of = r.headers.get("content-type", "").split(";")[0].strip()
    data = b"".join(chunks)
    if kind_of == "application/pdf" or data.startswith(b"%PDF-"):
        return data, "pdf"
    if kind_of in ("text/html", "application/xhtml+xml"):
        raise ValueError(f"{url} is a web page, not a file: give the link to the PDF or file")
    return data, _kind(httpx.URL(url).path)


def _read(path: Path, kind: Kind) -> tuple[list[str], str | None]:
    """A source's text page by page, and the title its file gives, if any."""
    if kind == "pdf":
        import pypdfium2 as pdfium

        doc = pdfium.PdfDocument(path)
        try:
            pages = [doc[i].get_textpage().get_text_range() for i in range(len(doc))]
            title = doc.get_metadata_dict().get("Title") or None
        finally:
            doc.close()
        return pages, title
    if kind == "pptx":
        from pptx import Presentation
        from pptx.shapes.autoshape import Shape

        deck = Presentation(str(path))
        slides = []
        for slide in deck.slides:
            text = [
                s.text_frame.text for s in slide.shapes if isinstance(s, Shape) and s.has_text_frame
            ]
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                text.append(slide.notes_slide.notes_text_frame.text)
            slides.append("\n".join(t for t in text if t))
        return slides, deck.core_properties.title or None
    if kind == "docx":
        from docx import Document

        doc = Document(str(path))
        return ["\n".join(p.text for p in doc.paragraphs)], doc.core_properties.title or None
    if kind == "ipynb":
        import nbformat

        notebook = nbformat.read(str(path), as_version=4)
        return [c.source for c in notebook.cells], None
    text = path.read_text(encoding="utf-8")
    heading = re.search(r"^# (.+)$", text, re.M) if kind == "md" else None
    return [text], heading.group(1).strip() if heading else None
