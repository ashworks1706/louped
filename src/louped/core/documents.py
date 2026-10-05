"""The text of PDFs, slide decks, Word documents and notebooks, read the same way wherever louped
reads one: sources/ for search and pins, reports/ for previews and louped check. Each reader
imports its library when called; they come with the sources extra, and notebook_html with the
notebooks extra."""

from __future__ import annotations

from pathlib import Path
from typing import NamedTuple


class Block(NamedTuple):
    """A paragraph of a Word document, or a table as text: a row a line, its cells joined by
    " | "."""

    text: str
    table: bool


def pdf_pages(path: Path) -> tuple[list[str], str | None]:
    """Each page's text, and the title the file gives."""
    import pypdfium2 as pdfium

    doc = pdfium.PdfDocument(path)
    try:
        pages = [doc[i].get_textpage().get_text_range() for i in range(len(doc))]
        title = doc.get_metadata_dict().get("Title") or None
    finally:
        doc.close()
    return pages, title


def pptx_slides(path: Path) -> tuple[list[str], str | None]:
    """Each slide's text, with its tables, grouped shapes and speaker notes, and the deck's
    title."""
    from pptx import Presentation

    deck = Presentation(str(path))
    slides = []
    for slide in deck.slides:
        text = [t for shape in slide.shapes for t in _shape_text(shape)]
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            text.append(slide.notes_slide.notes_text_frame.text)
        slides.append("\n".join(t for t in text if t))
    return slides, deck.core_properties.title or None


def _shape_text(shape: object) -> list[str]:
    from pptx.shapes.autoshape import Shape
    from pptx.shapes.graphfrm import GraphicFrame
    from pptx.shapes.group import GroupShape

    if isinstance(shape, GroupShape):
        return [t for s in shape.shapes for t in _shape_text(s)]
    if isinstance(shape, GraphicFrame) and shape.has_table:
        return [" | ".join(c.text for c in row.cells) for row in shape.table.rows]
    if isinstance(shape, Shape) and shape.has_text_frame:  # text boxes and placeholders
        return [shape.text_frame.text]
    return []


def docx_blocks(path: Path) -> tuple[list[Block], str | None]:
    """The document's paragraphs and tables in order, and its title."""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(path))
    blocks = []
    for block in doc.iter_inner_content():
        if isinstance(block, Paragraph):
            blocks.append(Block(block.text, table=False))
        elif isinstance(block, Table):
            rows = "\n".join(" | ".join(c.text for c in r.cells) for r in block.rows)
            blocks.append(Block(rows, table=True))
    return blocks, doc.core_properties.title or None


def notebook_cells(path: Path) -> list[str]:
    """Each cell's source."""
    import nbformat

    return [c.source for c in nbformat.read(str(path), as_version=4).cells]


def notebook_html(notebook: Path | bytes, dark: bool = False) -> str:
    """A notebook as one HTML page, its cells and saved outputs as JupyterLab shows them in its
    light or dark theme. The
    page's scripts (MathJax, interactive outputs) are for a frame that allows them; the app's does
    not, so math shows as its TeX and an interactive output as its static copy, if it saved one."""
    import nbformat
    from nbconvert import HTMLExporter

    try:
        text = notebook.read_text("utf-8") if isinstance(notebook, Path) else notebook.decode()
        node = nbformat.reads(text, 4)
    except (ValueError, nbformat.ValidationError) as exc:  # not UTF-8, not JSON, not a notebook
        raise ValueError(f"not a notebook louped can read: {exc}") from exc
    exporter = HTMLExporter(template_name="lab", theme="dark" if dark else "light")
    html, _ = exporter.from_notebook_node(node)
    return html
